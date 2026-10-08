import io

import pytest
from bson.binary import Binary
from docx import Document

from app.ingestion.chunker import chunk_text
from app.ingestion.parsers import ParseError, parse, supported_extensions
from app.ingestion.service import IngestionError, ingest_document

def words(text):
    return len(text.split())


def _pdf(pages):
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects)} 0 R >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_chunks_respect_token_budget_and_keep_paragraphs():
    paragraphs = [" ".join(f"p{i}w{j}" for j in range(40)) for i in range(10)]
    chunks = chunk_text("\n\n".join(paragraphs), words, max_tokens=100, overlap_tokens=20)
    assert all(words(c) <= 100 for c in chunks)
    assert len(chunks) == 5
    assert chunks[0] == paragraphs[0] + "\n\n" + paragraphs[1]


def test_long_paragraph_splits_on_sentences_with_overlap():
    sentences = [f"Sentence {i} " + " ".join(["word"] * 8) + "." for i in range(30)]
    chunks = chunk_text(" ".join(sentences), words, max_tokens=50, overlap_tokens=10)
    assert len(chunks) > 1
    assert all(words(c) <= 50 for c in chunks)
    assert chunks[1].startswith(chunks[0].split(". ")[-1].rstrip("."))


def test_oversized_sentence_is_split_on_words():
    chunks = chunk_text(" ".join(["token"] * 250), words, max_tokens=100, overlap_tokens=0)
    assert [words(c) for c in chunks] == [100, 100, 50]


def test_parse_text_markdown_and_unsupported():
    assert parse("notes.md", "# Title\n\nBody   text".encode())[0].text == "# Title\n\nBody text"
    assert parse("legacy.txt", "café".encode("cp1252"))[0].text == "café"
    assert ".pdf" in supported_extensions() and ".docx" in supported_extensions()
    with pytest.raises(ParseError, match="Unsupported file type"):
        parse("image.png", b"...")


def test_parse_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Vacation policy")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Years", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "0-3", "15"
    buffer = io.BytesIO()
    doc.save(buffer)
    [section] = parse("policy.docx", buffer.getvalue())
    assert section.text == "Vacation policy\n\nYears | Days\n0-3 | 15"


def test_parse_pdf_keeps_page_numbers():
    sections = parse("handbook.pdf", _pdf(["First page text", "Second page text"]))
    assert [(s.page, s.text) for s in sections] == [(1, "First page text"), (2, "Second page text")]


def test_ingest_stores_chunks_with_access_and_vectors(app, db):
    with app.app_context():
        doc = ingest_document("handbook.pdf", _pdf(["Vacation is 15 days", "Benefits start day one"]),
                              access=["group:hr", "user:admin"], uploaded_by="admin")
    chunks = list(db.chunks.find({"doc_id": doc["_id"]}).sort("chunk_index", 1))
    assert doc["title"] == "handbook" and doc["chunk_count"] == 2
    assert [c["page"] for c in chunks] == [1, 2]
    assert all(c["access"] == ["group:hr", "user:admin"] for c in chunks)
    assert isinstance(chunks[0]["embedding"], Binary)
    assert len(chunks[0]["embedding"].as_vector().data) == 384


def test_ingest_rejects_duplicates_empty_files_and_missing_access(app):
    with app.app_context():
        ingest_document("a.txt", b"hello world", access=["everyone"], uploaded_by="admin")
        with pytest.raises(IngestionError, match="already uploaded"):
            ingest_document("b.txt", b"hello world", access=["everyone"], uploaded_by="admin")
        with pytest.raises(IngestionError, match="No text"):
            ingest_document("empty.txt", b"   \n ", access=["everyone"], uploaded_by="admin")
        with pytest.raises(IngestionError, match="at least one group"):
            ingest_document("c.txt", b"other", access=[], uploaded_by="admin")


def test_ingestion_never_calls_the_llm(app, monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("LLM called during ingestion")

    monkeypatch.setattr(app.extensions["llm"], "chat", fail)
    with app.app_context():
        ingest_document("d.md", b"# Doc\n\nSome text", access=["everyone"], uploaded_by="admin")
